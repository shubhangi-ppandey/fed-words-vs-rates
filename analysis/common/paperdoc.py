"""Turn a research-paper markdown file into a formal .docx and .pdf (reusable across portfolio projects).

Supported markdown: '# ' title (first line), '## ' / '### ' headings, paragraphs, '- ' bullets, '1. ' numbered items,
pipe tables (first row = header) with an optional caption line 'Table: ...' directly above, images
'![Figure caption](path)', '**bold**' and '*italic*' inline, and '> ' block quotes (rendered as an abstract box).

Pipeline: python-docx builds the document (title page, table of contents field, numbered pages) ->
Microsoft Word (via pywin32, Windows only) updates the fields, saves and exports the PDF -> python-docx resets
the core properties (author, last modified by, title), so the files carry the author's name, not the machine user's.
"""
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def _field(paragraph, instr: str, placeholder: str = "") -> None:
    """Insert a Word field (e.g. 'TOC \\o "1-2" \\h \\z \\u' or 'PAGE')."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def _inline(paragraph, text: str, size=None, color=None) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        bold, italic = part.startswith("**"), part.startswith("*") and not part.startswith("**")
        run = paragraph.add_run(part.strip("*") if (bold or italic) else part)
        run.bold, run.italic = bold or None, italic or None
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def _styles(doc) -> None:
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.5
    for name, size in (("Heading 1", 14), ("Heading 2", 12.5), ("Title", 20)):
        s = doc.styles[name]
        s.font.name, s.font.size, s.font.bold = FONT, Pt(size), True
        s.font.color.rgb = NAVY
        s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _table(doc, rows: list, caption: str | None) -> None:
    if caption:
        p = doc.add_paragraph()
        _inline(p, caption, size=11)
        p.runs[0].bold = True
        p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in doc.styles] else "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        for j, cell in enumerate(r):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            _inline(p, cell.strip(), size=9.5)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def build_docx(md_path: Path, docx_path: Path, author: str, subtitle_lines: list) -> str:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    title = lines[0].lstrip("# ").strip()
    doc = Document()
    _styles(doc)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.5)

    # Title page
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    for line in subtitle_lines:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _inline(q, line, size=12.5)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_paragraph()  # not a Heading style, so it does not list itself in the TOC
    _inline(h, "**Contents**", size=14, color=NAVY)
    _field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', "(table of contents is generated when the PDF is built)")
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.footer.is_linked_to_previous = False  # no page number on the title/contents pages
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    body._sectPr.append(pg)
    footer = body.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(footer, "PAGE", "1")

    i, caption = 1, None
    para = []

    def flush():
        if para:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _inline(p, " ".join(para))
            para.clear()

    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            flush()
        elif ln.startswith("## "):
            flush()
            doc.add_paragraph(ln[3:].strip(), style="Heading 1")
        elif ln.startswith("### "):
            flush()
            doc.add_paragraph(ln[4:].strip(), style="Heading 2")
        elif ln.startswith("Table: "):
            flush()
            caption = ln[len("Table: "):].strip()
        elif ln.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"\s*:?-{2,}:?\s*", c) for c in cells):
                    rows.append(cells)
                i += 1
            _table(doc, rows, caption)
            caption = None
            continue
        elif m := re.match(r"!\[(.*)\]\((.*)\)", ln):
            flush()
            img = (md_path.parent / m.group(2)).resolve()
            doc.add_picture(str(img), width=Cm(15.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _inline(cp, m.group(1), size=10)
        elif ln.startswith("> "):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = p.paragraph_format.right_indent = Cm(1)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _inline(p, ln[2:], size=11)
        elif re.match(r"^\s*- ", ln):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            _inline(p, re.sub(r"^\s*- ", "", ln))
        elif re.match(r"^\s*\d+\. ", ln):
            flush()
            p = doc.add_paragraph(style="List Number")
            _inline(p, re.sub(r"^\s*\d+\. ", "", ln))
        else:
            para.append(ln.strip())
        i += 1
    flush()
    set_properties(doc, author, title)
    doc.save(docx_path)
    return title


def set_properties(doc, author: str, title: str) -> None:
    cp = doc.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.title = title
    cp.comments = ""
    cp.created = cp.modified = datetime.now()
    cp.revision = 1


def word_finalize(docx_path: Path, pdf_path: Path) -> None:
    """Update fields (TOC, page numbers) and export the PDF with Microsoft Word (Windows + pywin32)."""
    import win32com.client  # noqa: PLC0415

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible, word.DisplayAlerts = False, 0
    try:
        d = word.Documents.Open(str(docx_path.resolve()))
        d.Fields.Update()
        for toc in d.TablesOfContents:
            toc.Update()
        d.Save()
        d.ExportAsFixedFormat(OutputFileName=str(pdf_path.resolve()), ExportFormat=17,  # 17 = wdExportFormatPDF
                              OpenAfterExport=False, OptimizeFor=0, IncludeDocProps=True,
                              CreateBookmarks=1, DocStructureTags=True)
        d.Close(False)
    finally:
        word.Quit()


def make(md_path: Path, author: str, subtitle_lines: list) -> tuple[Path, Path]:
    docx_path, pdf_path = md_path.with_suffix(".docx"), md_path.with_suffix(".pdf")
    title = build_docx(md_path, docx_path, author, subtitle_lines)
    word_finalize(docx_path, pdf_path)
    d = Document(str(docx_path))  # Word stamped the machine user as last author: reset it
    set_properties(d, author, title)
    d.save(docx_path)
    return docx_path, pdf_path
